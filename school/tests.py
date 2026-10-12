import io
import json
import tempfile
from datetime import date, time
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import TestCase, override_settings, Client
from django.urls import reverse
from openpyxl import Workbook, load_workbook
from .models import *
from .policies import document_allowed
from .services import save_grade, success_rows, add_version, transfer, validate_upload

PASSWORD = 'OnlyForTests!Secure3479'

class WorkflowTests(TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.override = override_settings(MEDIA_ROOT=Path(self.temp.name)/'documents', EMAIL_BACKEND='django.core.mail.backends.locmem.EmailBackend')
        self.override.enable()
        self.addCleanup(self.override.disable)
        self.addCleanup(self.temp.cleanup)
        self.users = {}
        for role,_ in ROLES:
            self.users[role] = User.objects.create_user(username=role+'@school.test', email=role+'@school.test',
                password=PASSWORD, first_name=role, role=role, status='active', is_active=True)
        self.teacher = self.users['teacher']
        self.other = User.objects.create_user(username='other@school.test', email='other@school.test',password=PASSWORD,
            role='teacher',status='active')
        self.year = SchoolYear.objects.create(name='TEST 2026/27')
        self.period = Period.objects.create(year=self.year,name='Tremujori I',starts=date(2026,9,1),ends=date(2026,11,30))
        self.classroom = Classroom.objects.create(year=self.year,level=3,parallel='A',homeroom=self.users['homeroom'])
        self.otherclass = Classroom.objects.create(year=self.year,level=4,parallel='A')
        self.math = Subject.objects.create(name='Matematikë')
        self.language = Subject.objects.create(name='Gjuhë shqipe')
        self.assignment = Assignment.objects.create(classroom=self.classroom,subject=self.math,teacher=self.teacher)
        self.language_assignment = Assignment.objects.create(classroom=self.classroom,subject=self.language,teacher=self.other)
        self.other_assignment = Assignment.objects.create(classroom=self.otherclass,subject=self.math,teacher=self.other)
        self.test_type = AssessmentType.objects.create(year=self.year,name='Test',weight=2)
        self.oral_type = AssessmentType.objects.create(year=self.year,name='Me gojë',weight=1)
        self.student = Student.objects.create(first_name='TEST Arta',last_name='Shembull',registration_number='TEST-1')
        self.enrollment = Enrollment.objects.create(student=self.student,classroom=self.classroom,starts=date(2026,9,1))
        self.student2 = Student.objects.create(first_name='TEST Beni',last_name='Shembull',registration_number='TEST-2')
        self.enrollment2 = Enrollment.objects.create(student=self.student2,classroom=self.classroom,starts=date(2026,9,1))
        self.otherstudent = Student.objects.create(first_name='TEST Jashtë',last_name='Klase',registration_number='TEST-3')
        self.other_enrollment = Enrollment.objects.create(student=self.otherstudent,classroom=self.otherclass,starts=date(2026,9,1))

    def login(self, role): self.client.force_login(self.users[role])

    def grade(self,value=5,enrollment=None,assignment=None,kind=None):
        grade = Grade(enrollment=enrollment or self.enrollment,assignment=assignment or self.assignment,
            period=self.period,assessment_type=kind or self.test_type,value=value,date=date(2026,10,1))
        save_grade(grade,self.teacher)
        return grade

    def document(self,kind='shared',owner=None):
        doc = Document.objects.create(title='TEST Dokumenti',category='TEST',kind=kind,owner=owner or self.teacher,year=self.year)
        version = add_version(doc,SimpleUploadedFile('test.csv',b'emri,nota\nArta,5\nBeni,3\n'),doc.owner)
        return doc,version

    def test_registration_is_pending_and_cannot_self_promote(self):
        response = self.client.post(reverse('register'),{'first_name':'E re','last_name':'Punonjëse','email':'new@school.test',
            'position':'Drejtor','password1':PASSWORD,'password2':PASSWORD,'role':'admin','status':'active'})
        self.assertEqual(response.status_code,200)
        user = User.objects.get(email='new@school.test')
        self.assertEqual((user.role,user.status,user.is_active),('staff','pending',False))
        self.assertFalse(self.client.login(username=user.username,password=PASSWORD))

    def test_approval_and_blocking_revoke_existing_session(self):
        pending = User.objects.create_user(username='pending@school.test',email='pending@school.test',password=PASSWORD,status='pending',is_active=False)
        self.login('admin')
        self.assertEqual(self.client.post(reverse('approve_user',args=[pending.pk]),{'role':'teacher','status':'active'}).status_code,302)
        pending.refresh_from_db()
        worker = Client()
        self.assertTrue(worker.login(username=pending.username,password=PASSWORD))
        self.assertEqual(worker.get('/').status_code,200)
        self.client.post(reverse('approve_user',args=[pending.pk]),{'role':'teacher','status':'blocked'})
        self.assertEqual(worker.get('/').status_code,302)

    def test_last_admin_cannot_disable_self(self):
        self.login('admin')
        response = self.client.post(reverse('approve_user',args=[self.users['admin'].pk]),{'role':'staff','status':'disabled'})
        self.assertEqual(response.status_code,200)
        self.users['admin'].refresh_from_db()
        self.assertEqual(self.users['admin'].role,'admin')

    def test_all_role_menus_and_server_restrictions(self):
        from .policies import MODULES
        for role in self.users:
            self.login(role)
            self.assertEqual(self.client.get('/').status_code,200,role)
            for key,(_,model,read,write) in MODULES.items():
                response = self.client.get(reverse('record_list',args=[key]))
                self.assertEqual(response.status_code,200 if role in read else 403,(role,key))
                if role not in write:
                    self.assertEqual(self.client.get(reverse('record_create',args=[key])).status_code,403,(role,key,'create'))

    def test_teacher_student_scope(self):
        self.login('teacher')
        self.assertContains(self.client.get(reverse('record_list',args=['students'])),'TEST Arta')
        self.assertNotContains(self.client.get(reverse('record_list',args=['students'])),'TEST Jashtë')
        self.assertEqual(self.client.get(reverse('student',args=[self.otherstudent.pk])).status_code,404)

    def test_teacher_cannot_write_unassigned_grades(self):
        self.login('teacher')
        response = self.client.post(reverse('record_create',args=['grades']),{
            'enrollment':self.other_enrollment.pk,'assignment':self.other_assignment.pk,'period':self.period.pk,
            'assessment_type':self.test_type.pk,'value':5,'date':'2026-10-01','expected_revision':0})
        self.assertEqual(response.status_code,200)
        self.assertEqual(Grade.objects.count(),0)

    def test_teacher_grade_create_update_history_and_persistence(self):
        self.login('teacher')
        payload = {'enrollment':self.enrollment.pk,'assignment':self.assignment.pk,'period':self.period.pk,
            'assessment_type':self.test_type.pk,'value':5,'date':'2026-10-01','expected_revision':0}
        self.assertEqual(self.client.post(reverse('record_create',args=['grades']),payload).status_code,302)
        grade = Grade.objects.get()
        payload.update(value=4,expected_revision=grade.revision)
        self.assertEqual(self.client.post(reverse('record_edit',args=['grades',grade.pk]),payload).status_code,302)
        grade.refresh_from_db()
        self.assertEqual(grade.value,4)
        history = grade.gradehistory_set.first()
        self.assertEqual((history.previous['value'],history.current['value']),(5,4))
        self.client.post(reverse('logout'))
        self.assertTrue(self.client.login(username=self.teacher.username,password=PASSWORD))
        self.assertContains(self.client.get(reverse('grade_history',args=[grade.pk])),'Vlera e re: 4')

    def test_optimistic_grade_conflict(self):
        grade = self.grade()
        grade.value = 4
        save_grade(grade,self.teacher,1)
        grade.value = 3
        with self.assertRaises(ValidationError): save_grade(grade,self.teacher,1)
        self.assertEqual(Grade.objects.get(pk=grade.pk).value,4)

    def test_weighted_averages_incomplete_data_and_class_boundaries(self):
        self.grade(5)
        self.grade(2,kind=self.oral_type)
        self.grade(3,assignment=self.language_assignment)
        self.grade(1,enrollment=self.enrollment2)
        self.grade(1,enrollment=self.other_enrollment,assignment=self.other_assignment)
        row = success_rows([self.classroom],self.period)[0]
        self.assertEqual((row['total'],row['complete'],row['incomplete']),(2,1,1))
        self.assertEqual(row['individuals'][0]['average'],Decimal('3.50'))
        self.assertEqual(row['average'],Decimal('3.50'))
        self.assertEqual(row['subjects']['Matematikë']['average'],Decimal('2.50'))
        self.assertIsNone(row['individuals'][1]['average'])

    def test_existing_grades_keep_weight_snapshot(self):
        grade = self.grade()
        self.test_type.weight = 4
        self.test_type.save()
        grade.assessment_type = self.test_type
        grade.value = 4
        save_grade(grade,self.teacher,1)
        self.assertEqual(grade.weight,Decimal(2))

    def test_grade_class_year_date_validation(self):
        grade = self.grade()
        grade.enrollment = self.other_enrollment
        with self.assertRaises(ValidationError): grade.full_clean()
        grade.enrollment = self.enrollment
        grade.date = date(2027,1,1)
        with self.assertRaises(ValidationError): grade.full_clean()

    def test_archived_year_blocks_mutation(self):
        grade = self.grade()
        self.year.archived = True
        self.year.save()
        grade.period.year = self.year
        with self.assertRaises(ValidationError): save_grade(grade,self.teacher,1)
        self.login('admin')
        response = self.client.get(reverse('record_edit',args=['grades',grade.pk]))
        self.assertEqual(response.status_code,302)

    def test_multi_class_pedagogue_report_and_teacher_subject_scope(self):
        self.grade()
        self.grade(1,assignment=self.language_assignment)
        self.login('pedagogue')
        response = self.client.get(reverse('success'),{'period':self.period.pk})
        self.assertContains(response,'3-A')
        self.assertContains(response,'4-A')
        self.assertEqual(len(response.context['rows']),2)
        self.login('teacher')
        response = self.client.get(reverse('success'),{'period':self.period.pk})
        self.assertNotContains(response,'Gjuhë shqipe:')

    def test_month_and_year_reports_aggregate_only_selected_dates_and_year(self):
        self.grade(5)
        self.grade(3,assignment=self.language_assignment)
        self.login('pedagogue')
        for report_type,extra in [('month',{'month':'2026-10'}),('year',{}),('range',{'from':'2026-10-01','to':'2026-10-31'})]:
            response = self.client.get(reverse('success'),{'period':self.period.pk,'report_type':report_type,**extra})
            self.assertEqual(response.status_code,200)
            self.assertEqual(response.context['rows'][0]['average'],Decimal('4.00'))
        response = self.client.get(reverse('success'),{'period':self.period.pk,'report_type':'month','month':'2026-09'})
        self.assertIsNone(response.context['rows'][0]['average'])

    def test_document_confidentiality_including_administrator(self):
        doc,version = self.document('psychological',self.users['psychologist'])
        for role in ['admin','director','secretary','teacher','pedagogue','staff','homeroom']:
            self.login(role)
            self.assertEqual(self.client.get(reverse('document',args=[doc.pk])).status_code,404)
            self.assertEqual(self.client.get(reverse('download_document',args=[version.pk])).status_code,403)
            self.assertNotContains(self.client.get('/'),doc.title)
        grant = DocumentPermission(document=doc,user=self.users['director'],view=True,download=True)
        grant.full_clean()
        grant.save()
        self.login('director')
        self.assertEqual(self.client.get(reverse('document',args=[doc.pk])).status_code,200)
        self.assertEqual(self.client.get(reverse('download_document',args=[version.pk])).status_code,200)
        with self.assertRaises(ValidationError): DocumentPermission(document=doc,role='director').full_clean()

    def test_separate_document_view_download_permissions(self):
        doc,version = self.document()
        DocumentPermission.objects.create(document=doc,role='secretary',view=True,download=False)
        self.login('secretary')
        self.assertEqual(self.client.get(reverse('spreadsheet',args=[version.pk])).status_code,200)
        self.assertEqual(self.client.get(reverse('download_document',args=[version.pk])).status_code,403)
        self.assertEqual(self.client.get(reverse('spreadsheet',args=[version.pk]),{'format':'xlsx'}).status_code,403)

    def test_upload_download_version_and_review_notification(self):
        self.login('teacher')
        response = self.client.post(reverse('record_create',args=['documents']),{'title':'TEST Plan vjetor','category':'Planifikim',
            'kind':'pedagogical','year':self.year.pk,'file':SimpleUploadedFile('plan.pdf',b'%PDF-1.4\nTEST')})
        self.assertEqual(response.status_code,302)
        doc = Document.objects.get(title='TEST Plan vjetor')
        response = self.client.get(reverse('download_document',args=[doc.versions.first().pk]))
        self.assertEqual(b''.join(response.streaming_content),b'%PDF-1.4\nTEST')
        self.assertEqual(self.client.post(reverse('upload_version',args=[doc.pk]),{'file':SimpleUploadedFile('plan2.csv',b'a,b\n1,2')}).status_code,302)
        self.assertEqual(doc.versions.count(),2)
        self.client.post(reverse('submit_document',args=[doc.pk]))
        self.login('pedagogue')
        self.assertEqual(self.client.post(reverse('review_document',args=[doc.pk]),{'status':'changes','comment':'Plotësoni objektivat.'}).status_code,302)
        doc.refresh_from_db()
        self.assertEqual(doc.status,'changes')
        self.assertTrue(Notification.objects.filter(recipient=self.teacher,title__contains='status i ri').exists())

    def test_invalid_and_oversized_files(self):
        for filename,content in [('bad.pdf',b'not pdf'),('bad.xlsx',b'not zip'),('bad.exe',b'executable')]:
            with self.assertRaises(ValidationError): validate_upload(SimpleUploadedFile(filename,content))
        with self.assertRaises(ValidationError): validate_upload(SimpleUploadedFile('big.csv',b'x'*(20*1024*1024+1)))

    def test_spreadsheet_filter_sort_export_and_original_preservation(self):
        doc,version = self.document()
        self.login('teacher')
        response = self.client.get(reverse('spreadsheet',args=[version.pk]),{'q':'Arta','sort':'0'})
        self.assertContains(response,'Arta')
        self.assertNotContains(response,'Beni')
        response = self.client.get(reverse('spreadsheet',args=[version.pk]),{'q':'Arta','format':'xlsx'})
        sheet = load_workbook(io.BytesIO(response.content)).active
        self.assertEqual(sheet.cell(3,1).value,'Arta')
        with version.file.open('rb') as handle: self.assertIn(b'Beni',handle.read())

    def test_controlled_excel_import_preview_commit_and_duplicate_rejection(self):
        book = Workbook()
        book.active.append(['numri_regjistrit','emri','mbiemri','klasa','paralelja','data_regjistrimit'])
        book.active.append(['TEST-NEW','TEST Era','Importuar',3,'A','2026-09-01'])
        output = io.BytesIO()
        book.save(output)
        doc,_ = self.document(owner=self.users['secretary'])
        version = add_version(doc,SimpleUploadedFile('students.xlsx',output.getvalue()),doc.owner)
        self.login('secretary')
        url = reverse('import_students',args=[version.pk])
        self.assertContains(self.client.get(url),'TEST Era')
        self.assertFalse(Student.objects.filter(registration_number='TEST-NEW').exists())
        self.assertEqual(self.client.post(url,{'confirm':'yes'}).status_code,302)
        self.assertTrue(Enrollment.objects.filter(student__registration_number='TEST-NEW').exists())
        self.assertContains(self.client.post(url,{'confirm':'yes'}),'dyfishuar')
        self.assertEqual(Student.objects.filter(registration_number='TEST-NEW').count(),1)

    def test_exports_pdf_excel_csv(self):
        self.grade()
        self.login('pedagogue')
        for fmt in ['pdf','xlsx','csv']:
            response = self.client.get(reverse('success'),{'period':self.period.pk,'format':fmt})
            self.assertEqual(response.status_code,200,fmt)
            if fmt == 'pdf': self.assertTrue(response.content.startswith(b'%PDF'))
            elif fmt == 'xlsx': self.assertEqual(load_workbook(io.BytesIO(response.content)).active.cell(2,1).value,'Klasa')
            else: self.assertIn('TEST 2026/27',response.content.decode('utf-8-sig'))

    def test_export_neutralizes_spreadsheet_formulas(self):
        self.student.first_name = '=HYPERLINK("https://example.invalid")'
        self.student.save()
        self.login('secretary')
        response = self.client.get(reverse('record_list',args=['students']),{'format':'xlsx'})
        cells = [cell for row in load_workbook(io.BytesIO(response.content)).active for cell in row]
        self.assertFalse(any(cell.data_type=='f' for cell in cells))

    def test_schedule_conflicts_for_teacher_classroom_and_room(self):
        Schedule.objects.create(assignment=self.assignment,weekday=0,starts=time(9),ends=time(10),room='A1')
        for assignment,room in [(self.language_assignment,'B1'),(self.other_assignment,'A1')]:
            with self.assertRaises(ValidationError): Schedule(assignment=assignment,weekday=0,starts=time(9,30),ends=time(10,30),room=room).full_clean()
        Schedule(assignment=self.other_assignment,weekday=0,starts=time(10),ends=time(11),room='A1').full_clean()

    def test_transfer_preserves_prior_grades_and_enrollment(self):
        grade = self.grade()
        new = transfer(self.enrollment,self.otherclass,date(2026,10,2),self.users['secretary'])
        self.enrollment.refresh_from_db()
        self.assertEqual(self.enrollment.ends,date(2026,10,1))
        self.assertEqual(Grade.objects.get(pk=grade.pk).enrollment_id,self.enrollment.pk)
        self.assertEqual(new.classroom_id,self.otherclass.pk)

    def test_duplicate_enrollment_in_same_year(self):
        with self.assertRaises(ValidationError): Enrollment(student=self.student,classroom=self.otherclass,starts=date(2026,9,1)).full_clean()

    def test_confidential_cases_not_available_to_other_psychologist(self):
        case = PsychologyCase.objects.create(student=self.student,psychologist=self.users['psychologist'],title='TEST Privat',notes='Konfidenciale')
        self.login('psychologist')
        self.assertContains(self.client.get(reverse('record_list',args=['cases'])),'TEST Privat')
        self.login('director')
        self.assertEqual(self.client.get(reverse('record_list',args=['cases'])).status_code,403)
        self.assertEqual(self.client.get(reverse('case',args=[case.pk])).status_code,404)
        case.authorized.add(self.users['director'])
        self.assertContains(self.client.get(reverse('case',args=[case.pk])),'Konfidenciale')
        self.assertNotContains(self.client.get('/'),'Konfidenciale')
        ModulePermission.objects.create(role='director',module='cases',can_view=False,can_edit=False)
        self.assertEqual(self.client.get(reverse('case',args=[case.pk])).status_code,403)

    def test_admin_creation_and_reports(self):
        self.login('admin')
        response = self.client.post(reverse('create_user'),{'first_name':'TEST I ri','last_name':'Punonjës','email':'created@school.test',
            'password1':PASSWORD,'password2':PASSWORD,'role':'secretary','status':'active'})
        self.assertEqual(response.status_code,302)
        user = User.objects.get(email='created@school.test')
        self.assertEqual((user.role,user.status,user.is_active),('secretary','active',True))
        self.assertEqual(self.client.get(reverse('reports')).status_code,200)
        self.assertContains(self.client.get(reverse('usage')),'Krijim llogarie')
        self.login('teacher')
        self.assertEqual(self.client.get(reverse('create_user')).status_code,403)
        self.assertEqual(self.client.get(reverse('usage')).status_code,403)

    def test_document_deadline_notification_is_private_and_repeatable(self):
        from django.utils import timezone
        doc,version = self.document('pedagogical')
        doc.due_date = timezone.localdate()
        doc.save()
        call_command('notify_deadlines',stdout=io.StringIO())
        call_command('notify_deadlines',stdout=io.StringIO())
        notes = Notification.objects.filter(title='Afati i dokumentacionit pedagogjik')
        self.assertEqual(notes.count(),1)
        self.assertEqual(notes.get().recipient,self.teacher)

    def test_notifications_are_recipient_scoped(self):
        notification = Notification.objects.create(recipient=self.users['psychologist'],title='TEST Privat',body='Mesazh privat')
        self.login('teacher')
        self.assertNotContains(self.client.get(reverse('notifications')),'TEST Privat')
        self.client.post(reverse('notifications'),{'id':notification.pk})
        notification.refresh_from_db()
        self.assertFalse(notification.read)

    def test_csrf_required_on_mutations(self):
        client = Client(enforce_csrf_checks=True)
        client.force_login(self.users['admin'])
        self.assertEqual(client.post(reverse('approve_user',args=[self.teacher.pk]),{'role':'admin','status':'active'}).status_code,403)

    def test_login_rate_limit(self):
        for _ in range(8): self.assertEqual(self.client.post(reverse('login'),{'username':'unknown@school.test','password':'wrong'}).status_code,200)
        self.assertEqual(self.client.post(reverse('login'),{'username':'unknown@school.test','password':'wrong'}).status_code,429)

    def test_password_reset_email_for_active_account(self):
        from django.core import mail
        response = self.client.post(reverse('password_reset'),{'email':self.teacher.email})
        self.assertEqual(response.status_code,302)
        self.assertEqual(len(mail.outbox),1)
        self.assertIn('/fjalekalimi/',mail.outbox[0].body)

    def test_no_public_media_route(self):
        doc,version = self.document('psychological',self.users['psychologist'])
        self.assertEqual(self.client.get('/media/'+version.file.name).status_code,404)

    def test_module_permission_restriction_applies_to_direct_routes(self):
        doc,version = self.document()
        rule = ModulePermission(role='teacher',module='documents',can_view=False,can_edit=False)
        rule.full_clean()
        rule.save()
        self.login('teacher')
        self.assertEqual(self.client.get(reverse('record_list',args=['documents'])).status_code,403)
        self.assertEqual(self.client.get(reverse('download_document',args=[version.pk])).status_code,403)
        self.assertNotContains(self.client.get('/'),'href="/modulet/documents/"')

    def test_spreadsheet_edit_creates_new_version_and_rejects_stale(self):
        doc,version = self.document()
        self.login('teacher')
        url = reverse('edit_spreadsheet',args=[version.pk])
        self.assertEqual(self.client.get(url).status_code,200)
        payload = {'cell_0_0':'Arta','cell_0_1':'4','cell_1_0':'Beni','cell_1_1':'3'}
        self.assertEqual(self.client.post(url,payload).status_code,302)
        self.assertEqual(doc.versions.count(),2)
        sheet = load_workbook(doc.versions.first().file.open('rb')).active
        self.assertEqual(sheet.cell(2,2).value,'4')
        self.assertContains(self.client.post(url,payload),'version më i ri')
        self.assertEqual(doc.versions.count(),2)

    def test_grade_and_attendance_import_validation_and_history(self):
        doc,_ = self.document()
        grade_file = SimpleUploadedFile('grades.csv',b'numri_regjistrit,klasa,paralelja,lenda,periudha,lloji,data,nota\nTEST-1,3,A,Matematik\xc3\xab,Tremujori I,Test,2026-10-01,4\n')
        version = add_version(doc,grade_file,self.teacher)
        self.login('teacher')
        url = reverse('import_academic',args=[version.pk,'grades'])
        response = self.client.get(url)
        self.assertEqual(response.status_code,200)
        self.assertEqual(len(response.context['records']),1)
        self.assertEqual(response.context['errors'],[])
        self.assertEqual(self.client.post(url,{'confirm':'yes'}).status_code,302)
        self.assertEqual(Grade.objects.count(),1)
        self.assertEqual(GradeHistory.objects.count(),1)
        self.assertContains(self.client.post(url,{'confirm':'yes'}),'importuar tashmë')
        attendance_file = SimpleUploadedFile('attendance.csv',b'numri_regjistrit,klasa,paralelja,data,statusi,arsyetimi\nTEST-1,3,A,2026-10-02,absent,\n')
        version = add_version(doc,attendance_file,self.teacher)
        url = reverse('import_academic',args=[version.pk,'attendance'])
        self.assertEqual(self.client.post(url,{'confirm':'yes'}).status_code,302)
        self.assertEqual(Attendance.objects.get().status,'absent')

    def test_import_academic_rejects_unassigned_class_without_partial_writes(self):
        doc,_ = self.document()
        version = add_version(doc,SimpleUploadedFile('bad.csv',b'numri_regjistrit,klasa,paralelja,data,statusi,arsyetimi\nTEST-1,3,A,2026-10-02,present,\nTEST-3,4,A,2026-10-02,present,\n'),self.teacher)
        self.login('teacher')
        response = self.client.post(reverse('import_academic',args=[version.pk,'attendance']),{'confirm':'yes'})
        self.assertEqual(response.status_code,200)
        self.assertTrue(response.context['errors'])
        self.assertEqual(Attendance.objects.count(),0)

    def test_public_auth_and_reset_templates_render(self):
        for url in ['login','register','password_reset','password_reset_done','password_reset_complete']:
            self.assertEqual(self.client.get(reverse(url)).status_code,200,url)

    def test_mutable_module_forms_render(self):
        from .policies import MODULES
        for key,(_,model,read,write) in MODULES.items():
            if not write: continue
            role = 'admin' if 'admin' in write else 'psychologist'
            self.login(role)
            self.assertEqual(self.client.get(reverse('record_create',args=[key])).status_code,200,key)

class BootstrapTests(TestCase):
    def test_bootstrap_has_no_default_password_and_is_one_time(self):
        with patch('school.management.commands.bootstrap_admin.getpass',return_value=PASSWORD):
            call_command('bootstrap_admin',email='owner@school.test',stdout=io.StringIO())
        admin = User.objects.get()
        self.assertTrue(admin.check_password(PASSWORD))
        self.assertEqual((admin.role,admin.status),('admin','active'))
        with self.assertRaises(CommandError): call_command('bootstrap_admin',email='second@school.test')

    def test_bootstrap_rejects_weak_password(self):
        with patch('school.management.commands.bootstrap_admin.getpass',return_value='1234'):
            with self.assertRaises(CommandError): call_command('bootstrap_admin',email='owner@school.test')
        self.assertFalse(User.objects.exists())
