import csv
import hashlib
import io
import json
from datetime import timedelta
from decimal import Decimal
from functools import wraps
from pathlib import Path
from django import forms
from django.contrib import messages
from django.contrib.auth import views as auth_views
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied, ValidationError, ObjectDoesNotExist, MultipleObjectsReturned
from django.db import transaction, IntegrityError
from django.db.models import Q, Count, F
from django.http import FileResponse, HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_POST
from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from xml.sax.saxutils import escape
from .models import *
from .policies import *
from .forms import *
from .services import *

def access(*roles):
    def decorator(view):
        @login_required
        @wraps(view)
        def wrapped(request, *args, **kwargs):
            if request.user.role not in roles: raise PermissionDenied
            return view(request, *args, **kwargs)
        return wrapped
    return decorator

class LoginView(auth_views.LoginView):
    template_name = 'school/login.html'
    authentication_form = LoginForm
    def post(self, request, *args, **kwargs):
        identity = request.POST.get('username', '').lower().strip()
        # Rate limit identity and transport IP separately; no raw address/email in this table.
        keys = [hashlib.sha256(value.encode()).hexdigest() for value in
                ['identity:' + identity, 'ip:' + request.META.get('REMOTE_ADDR', '')]]
        self.attempt_keys = keys
        recent = timezone.now() - timedelta(minutes=15)
        if LoginAttempt.objects.filter(key__in=keys, failures__gte=8, updated__gte=recent).exists():
            return HttpResponse('Shumë përpjekje. Provoni pas 15 minutash.', status=429)
        LoginAttempt.objects.filter(key__in=keys, updated__lt=recent).update(failures=0)
        return super().post(request, *args, **kwargs)
    def form_invalid(self, form):
        for key in self.attempt_keys:
            attempt, _ = LoginAttempt.objects.get_or_create(key=key)
            LoginAttempt.objects.filter(pk=attempt.pk).update(failures=F('failures') + 1, updated=timezone.now())
        audit(None, 'Hyrje e pasuksesshme', 'authentication')
        return super().form_invalid(form)
    def form_valid(self, form):
        user = form.get_user()
        if user.status != 'active': raise PermissionDenied
        LoginAttempt.objects.filter(key__in=self.attempt_keys).delete()
        audit(user, 'Hyrje në sistem', user.pk)
        return super().form_valid(form)

def register(request):
    form = RegistrationForm(request.POST or None)
    if request.method == 'POST' and form.is_valid():
        with transaction.atomic():
            user = form.save()
            audit(None, 'Kërkesë për llogari', user.pk)
            for admin in User.objects.filter(role='admin', status='active'):
                Notification.objects.create(recipient=admin, title='Kërkesë e re për llogari', body=f'{user} kërkon miratim.')
        return render(request, 'school/message.html', {'title': 'Në pritje të miratimit',
            'body': 'Kërkesa juaj u ruajt. Administratori duhet ta verifikojë dhe ta aktivizojë para hyrjes.'})
    return render(request, 'school/form.html', {'form': form, 'title': 'Kërkesë për llogari punonjësi'})

@login_required
def dashboard(request):
    user = request.user
    cards = []
    for module, label in [('students','Nxënës'), ('classes','Klasa'), ('documents','Dokumente të qasshme')]:
        if user.role in MODULES[module][2]: cards.append((label, scoped(user, module).count()))
    if user.role in {'admin','director','secretary'}: cards.append(('Punonjës aktivë', User.objects.filter(status='active').count()))
    if user.role in MODULES['attendance'][2]:
        cards.append(('Mungesa të paarsyetuara', scoped(user, 'attendance').filter(status='absent').count()))
    if user.role in {'admin','director','pedagogue'}:
        cards.append(('Dokumente në shqyrtim', documents_for(user).filter(kind='pedagogical', status__in=['submitted','review','changes']).count()))
    return render(request, 'school/dashboard.html', {'cards': cards,
        'notifications': Notification.objects.filter(recipient=user)[:6],
        'events': scoped(user, 'events').filter(ends__gte=timezone.now()).order_by('starts')[:5]})

@access('admin')
def users(request):
    queryset = User.objects.all().order_by('status', 'last_name')
    if request.GET.get('status'): queryset = queryset.filter(status=request.GET['status'])
    return render(request, 'school/users.html', {'users': queryset})

@access('admin')
def create_user(request):
    form = AdminCreationForm(request.POST or None)
    if request.method == 'POST' and form.is_valid():
        with transaction.atomic():
            user = form.save()
            audit(request.user,'Krijim llogarie nga administratori',user.pk,{'role':user.role,'status':user.status})
        messages.success(request,'Llogaria u krijua. Dorëzoni kredencialet përmes një kanali të sigurt.')
        return redirect('users')
    return render(request,'school/form.html',{'title':'Krijo llogari punonjësi','form':form})

@login_required
def reports(request):
    return render(request,'school/reports.html',{'reports':[(key,value[0]) for key,value in MODULES.items()
        if allowed(request.user,key) and key not in {'settings','permissions'}]})

@access('admin')
def usage_report(request):
    rows = list(Audit.objects.values('actor__first_name','actor__last_name','actor__email','action').annotate(count=Count('id')).order_by('actor__email','action'))
    headers = ['Punonjësi','Email-i','Veprimi','Numri']
    data = [[f"{row['actor__first_name'] or ''} {row['actor__last_name'] or ''}",row['actor__email'] or 'Pa hyrje',row['action'],row['count']] for row in rows]
    if request.GET.get('format'): return export_table(request,'Raporti i përdorimit të aplikacionit',headers,data)
    return render(request,'school/usage.html',{'rows':data})

@access('admin')
def approve_user(request, pk):
    user = get_object_or_404(User, pk=pk)
    form = UserApprovalForm(request.POST or None, instance=user)
    if request.method == 'POST' and form.is_valid():
        with transaction.atomic():
            # Serialize account lifecycle to preserve a last administrator.
            list(User.objects.select_for_update().filter(role='admin').values_list('id', flat=True))
            updated = form.save(commit=False)
            if updated.pk == request.user.pk and (updated.role != 'admin' or updated.status != 'active'):
                form.add_error(None, 'Nuk mund të hiqni qasjen tuaj administrative.')
            else:
                updated.is_active = updated.status == 'active'
                updated.is_staff = False
                updated.save()
                audit(request.user, 'Ndryshim roli / statusi', updated.pk, {'role': updated.role, 'status': updated.status})
                if updated.is_active:
                    Notification.objects.create(recipient=updated, title='Llogaria juaj është aktive', body='Mund të hyni në panelin tuaj.')
                return redirect('users')
    return render(request, 'school/form.html', {'form': form, 'title': f'Autorizimet: {user}'})

def list_columns(module):
    return {
        'students': ['first_name','last_name','registration_number'],
        'enrollments': ['student','classroom','starts','ends'],
        'classes': ['level','parallel','year','homeroom'],
        'subjects': ['name'], 'assignments': ['classroom','subject','teacher'],
        'years': ['name','archived','average_method','pass_threshold'],
        'periods': ['year','name','starts','ends'], 'types': ['year','name','weight'],
        'grades': ['enrollment','assignment','period','value','date'],
        'attendance': ['enrollment','date','status','reason'],
        'documents': ['title','category','kind','owner','year','status'],
        'cases': ['title','student','next_meeting','closed'],
        'schedule': ['assignment','weekday','starts','ends','room'],
        'events': ['title','starts','ends'], 'protocol': ['number','title','direction','date','status'],
        'audit': ['created_at','actor','action','target'], 'settings': ['name','retention_years'],
        'permissions': ['role','module','can_view','can_edit'],
    }[module]

def text_value(obj, name):
    display = getattr(obj, f'get_{name}_display', None)
    value = display() if display else getattr(obj, name)
    if value is None: return '—'
    if value is True: return 'Po'
    if value is False: return 'Jo'
    return str(value)

def filtered_records(request, module):
    qs = scoped(request.user, module)
    query = request.GET.get('q', '').strip()
    if query:
        search = Q()
        for f in MODULES[module][1]._meta.fields:
            if isinstance(f, (models.CharField, models.TextField)): search |= Q(**{f.name+'__icontains': query})
        if module in ('grades', 'attendance', 'enrollments'):
            prefix = 'enrollment__student__' if module != 'enrollments' else 'student__'
            search |= Q(**{prefix+'first_name__icontains': query}) | Q(**{prefix+'last_name__icontains': query})
        qs = qs.filter(search)
    year, classroom = request.GET.get('year'), request.GET.get('classroom')
    if year:
        key = {'students':'enrollment__classroom__year', 'enrollments':'classroom__year',
               'grades':'period__year', 'attendance':'enrollment__classroom__year',
               'schedule':'assignment__classroom__year', 'assignments':'classroom__year'}.get(module, 'year')
        if module in {'students','enrollments','grades','attendance','schedule','assignments','classes','periods','types','documents'}:
            qs = qs.filter(**{key: year})
    if classroom:
        key = {'students':'enrollment__classroom','enrollments':'classroom','grades':'assignment__classroom',
               'attendance':'enrollment__classroom','schedule':'assignment__classroom','assignments':'classroom','documents':'classroom'}.get(module)
        if key: qs = qs.filter(**{key: classroom})
    for key in ('status', 'category', 'kind'):
        if request.GET.get(key) and key in [f.name for f in MODULES[module][1]._meta.fields]: qs = qs.filter(**{key: request.GET[key]})
    if module == 'documents':
        if request.GET.get('owner'): qs = qs.filter(owner_id=request.GET['owner'])
        if request.GET.get('subject'): qs = qs.filter(subject_id=request.GET['subject'])
    if module == 'grades':
        if request.GET.get('subject'): qs = qs.filter(assignment__subject_id=request.GET['subject'])
        if request.GET.get('period'): qs = qs.filter(period_id=request.GET['period'])
    return qs.distinct()

@login_required
def record_list(request, module):
    if module not in MODULES or not allowed(request.user,module): raise PermissionDenied
    columns = list_columns(module)
    qs = filtered_records(request, module)
    if request.GET.get('format'):
        if module == 'documents':
            qs = [obj for obj in qs if document_allowed(request.user, obj, 'download')]
        return export_table(request, MODULES[module][0],
            [str(MODULES[module][1]._meta.get_field(c).verbose_name) for c in columns],
            [[text_value(obj,c) for c in columns] for obj in qs])
    from django.core.paginator import Paginator
    page = Paginator(qs.order_by('-created_at') if hasattr(qs.model, 'created_at') else qs.order_by('pk'), 30).get_page(request.GET.get('page'))
    rows = [{'object': obj, 'values': [text_value(obj,c) for c in columns], 'editable': writable(request.user,module,obj)} for obj in page]
    return render(request, 'school/list.html', {'title': MODULES[module][0], 'module': module, 'rows': rows,
        'headers': [str(MODULES[module][1]._meta.get_field(c).verbose_name) for c in columns],
        'can_create': writable(request.user,module) and (module != 'settings' or not SchoolSettings.objects.exists()), 'page': page,
        'years': SchoolYear.objects.all(), 'classrooms': classes_for(request.user),
        'subjects': Subject.objects.all(), 'periods': Period.objects.all(),
        'owners': User.objects.filter(status='active') if request.user.role in {'admin','director','pedagogue'} else User.objects.filter(pk=request.user.pk),
        'export_query': request.GET.urlencode()})

@login_required
def record_edit(request, module, pk=None):
    if module not in MODULES or not allowed(request.user,module): raise PermissionDenied
    if module == 'settings' and not pk and SchoolSettings.objects.exists():
        return redirect('record_edit',module='settings',pk=SchoolSettings.objects.first().pk)
    obj = get_object_or_404(scoped(request.user, module), pk=pk) if pk else None
    if not writable(request.user, module, obj): raise PermissionDenied
    if obj and year_of(obj) and year_of(obj).archived:
        messages.error(request, 'Viti është arkivuar; të dhënat ruhen vetëm për lexim.')
        return redirect('record_list', module=module)
    form = record_form(module, request.user, request.POST or None, request.FILES or None, obj)
    upload = UploadForm(request.POST or None, request.FILES or None) if module == 'documents' and not obj else None
    if request.method == 'POST' and form.is_valid() and (not upload or upload.is_valid()):
        try:
            with transaction.atomic():
                item = form.save(commit=False)
                if year_of(item) and year_of(item).archived: raise ValidationError('Viti është arkivuar.')
                if module == 'grades': save_grade(item, request.user, form.cleaned_data['expected_revision'])
                else:
                    if hasattr(item, 'owner_id') and not obj: item.owner = request.user
                    if hasattr(item, 'author_id'): item.author = request.user
                    if module == 'enrollments' and obj:
                        original = Enrollment.objects.get(pk=obj.pk)
                        if original.student_id != item.student_id or original.classroom_id != item.classroom_id:
                            raise ValidationError('Përdorni transferimin për të ndryshuar klasën. Nxënësi i regjistrimit nuk ndryshohet.')
                    item.full_clean()
                    item.save()
                    form.save_m2m()
                    audit(request.user, f'Ruajtje: {module}', item.pk)
                    if module == 'schedule':
                        recipients = User.objects.filter(Q(pk=item.assignment.teacher_id) | Q(pk=item.assignment.classroom.homeroom_id),status='active')
                        for person in recipients:
                            Notification.objects.create(recipient=person,title='Ndryshim në orar',body=str(item))
                    if upload: add_version(item, upload.cleaned_data['file'], request.user)
                    if module == 'attendance' and item.status == 'absent':
                        count = Attendance.objects.filter(enrollment=item.enrollment, status='absent').count()
                        if count % 5 == 0:
                            recipients = User.objects.filter(Q(role__in=['director','pedagogue']) | Q(pk=item.enrollment.classroom.homeroom_id), status='active')
                            for person in recipients:
                                Notification.objects.create(recipient=person, title='Mungesa të përsëritura', body=f'{item.enrollment.student}: {count} mungesa të paarsyetuara.')
                messages.success(request, 'Të dhënat u ruajtën në databazë.')
                if module == 'documents': return redirect('document', pk=item.pk)
                return redirect('record_list', module=module)
        except (ValidationError, IntegrityError) as error:
            form.add_error(None, error if isinstance(error, ValidationError) else 'Konflikt të dhënash. Rifreskoni dhe provoni përsëri.')
    return render(request, 'school/form.html', {'form': form, 'upload_form': upload,
        'title': ('Ndrysho: ' if obj else 'Shto: ') + MODULES[module][0]})

@access('admin','secretary')
def transfer_student(request, pk):
    enrollment = get_object_or_404(Enrollment, pk=pk)
    form = TransferForm(request.POST or None)
    form.fields['destination'].queryset = Classroom.objects.filter(year=enrollment.classroom.year, year__archived=False).exclude(pk=enrollment.classroom_id)
    if request.method == 'POST' and form.is_valid():
        try:
            transfer(enrollment, form.cleaned_data['destination'], form.cleaned_data['date'], request.user)
            messages.success(request, 'Transferimi u ruajt; historiku i klasës së mëparshme u ruajt.')
            return redirect('record_list', module='enrollments')
        except ValidationError as error: form.add_error(None,error)
    return render(request, 'school/form.html', {'title': f'Transfero: {enrollment.student}', 'form': form})

@login_required
def student_detail(request, pk):
    if request.user.role not in MODULES['students'][2]: raise PermissionDenied
    student = get_object_or_404(scoped(request.user,'students'), pk=pk)
    grades = scoped(request.user,'grades').filter(enrollment__student=student) if request.user.role in MODULES['grades'][2] else []
    attendance = scoped(request.user,'attendance').filter(enrollment__student=student) if request.user.role in MODULES['attendance'][2] else []
    return render(request, 'school/student.html', {'student': student, 'grades': grades,
        'enrollments': scoped(request.user,'enrollments').filter(student=student), 'attendance': attendance,
        'can_success': request.user.role in MODULES['grades'][2]})

@login_required
def grade_history(request, pk):
    if request.user.role not in MODULES['grades'][2]: raise PermissionDenied
    grade = get_object_or_404(scoped(request.user,'grades'), pk=pk)
    return render(request, 'school/history.html', {'grade': grade, 'history': grade.gradehistory_set.select_related('actor')})

@login_required
def authorized_cases(request, pk=None):
    if ModulePermission.objects.filter(role=request.user.role,module='cases',can_view=False).exists(): raise PermissionDenied
    cases = PsychologyCase.objects.filter(Q(psychologist=request.user) | Q(authorized=request.user)).distinct()
    case = get_object_or_404(cases,pk=pk) if pk else None
    return render(request,'school/cases.html',{'cases':cases,'case':case})

@login_required
def document_detail(request, pk):
    document = get_object_or_404(documents_for(request.user), pk=pk)
    edit = document_allowed(request.user,document,'edit') and not document.year.archived
    review = document.kind == 'pedagogical' and request.user.role in {'admin','director','pedagogue'} and not document.year.archived
    return render(request, 'school/document.html', {'document': document, 'versions': document.versions.all(),
        'comments': document.comments.select_related('author'), 'upload_form': UploadForm(), 'review_form': ReviewForm(),
        'can_edit': edit, 'can_review': review, 'can_manage': document.owner_id == request.user.pk or (request.user.role == 'admin' and document.kind != 'psychological'),
        'can_delete': document_allowed(request.user,document,'delete') and not document.year.archived,
        'can_submit': document.owner_id == request.user.pk and document.kind == 'pedagogical' and document.status in {'draft','changes'} and not document.year.archived,
        'can_download': document_allowed(request.user,document,'download'), 'permissions': document.access.all()})

@login_required
@require_POST
def upload_version(request, pk):
    document = get_object_or_404(documents_for(request.user), pk=pk)
    if not document_allowed(request.user,document,'edit') or document.year.archived: raise PermissionDenied
    form = UploadForm(request.POST, request.FILES)
    if form.is_valid():
        add_version(document, form.cleaned_data['file'], request.user)
        # Updated content always needs a new review.
        document.status = 'draft'
        document.save(update_fields=['status'])
        messages.success(request, 'Versioni i ri u ruajt si draft.')
    else: messages.error(request, str(form.errors.as_text()))
    return redirect('document', pk=pk)

@login_required
def download_document(request, pk):
    version = get_object_or_404(DocumentVersion.objects.select_related('document'), pk=pk)
    if not document_allowed(request.user, version.document, 'download'): raise PermissionDenied
    audit(request.user,'Shkarkim dokumenti', version.document_id, {'version':version.number})
    response = FileResponse(version.file.open('rb'), as_attachment=True, filename=version.original_name)
    response['X-Content-Type-Options'] = 'nosniff'
    return response

@login_required
@require_POST
def submit_document(request, pk):
    doc = get_object_or_404(documents_for(request.user), pk=pk)
    if doc.owner_id != request.user.pk or doc.year.archived or doc.kind != 'pedagogical': raise PermissionDenied
    if doc.status not in {'draft','changes'}: raise PermissionDenied
    doc.status = 'submitted'
    doc.save(update_fields=['status'])
    audit(request.user,'Dorëzim dokumenti',doc.pk)
    for user in User.objects.filter(role__in=['director','pedagogue'], status='active'):
        Notification.objects.create(recipient=user,title='Dokument pedagogjik për shqyrtim',body=doc.title)
    return redirect('document', pk=pk)

@access('admin','director','pedagogue')
@require_POST
def review_document(request, pk):
    doc = get_object_or_404(documents_for(request.user), pk=pk, kind='pedagogical')
    if doc.year.archived: raise PermissionDenied
    form = ReviewForm(request.POST)
    if form.is_valid():
        with transaction.atomic():
            doc.status = form.cleaned_data['status']
            doc.save(update_fields=['status'])
            DocumentComment.objects.create(document=doc,author=request.user,text=form.cleaned_data['comment'],status=doc.status)
            Notification.objects.create(recipient=doc.owner,title='Dokumentacioni pedagogjik: status i ri',body=f'{doc.title}: {doc.get_status_display()}')
            audit(request.user,'Shqyrtim dokumenti',doc.pk,{'status':doc.status})
    else: messages.error(request,str(form.errors.as_text()))
    return redirect('document',pk=pk)

@login_required
def document_permission(request, pk):
    doc = get_object_or_404(documents_for(request.user),pk=pk)
    if not (doc.owner_id == request.user.pk or (request.user.role == 'admin' and doc.kind != 'psychological')) or doc.year.archived: raise PermissionDenied
    cls = forms.modelform_factory(DocumentPermission,fields=['role','user','view','download','edit','delete'])
    form = cls(request.POST or None,instance=DocumentPermission(document=doc))
    form.fields['user'].queryset = User.objects.filter(status='active',is_active=True)
    if request.method == 'POST' and form.is_valid():
        permission = form.save()
        audit(request.user,'Autorizim dokumenti',doc.pk,{'permission':str(permission.pk)})
        return redirect('document',pk=pk)
    return render(request,'school/form.html',{'title':f'Lejet: {doc.title}','form':form})

@login_required
@require_POST
def revoke_permission(request, pk):
    permission = get_object_or_404(DocumentPermission,pk=pk)
    doc = permission.document
    if not (doc.owner_id == request.user.pk or (request.user.role == 'admin' and doc.kind != 'psychological')) or doc.year.archived: raise PermissionDenied
    audit(request.user,'Heqje autorizimi dokumenti',doc.pk,{'permission':str(permission.pk)})
    permission.delete()
    return redirect('document',pk=doc.pk)

@login_required
@require_POST
def delete_document(request, pk):
    doc = get_object_or_404(documents_for(request.user),pk=pk)
    if not document_allowed(request.user,doc,'delete') or doc.year.archived: raise PermissionDenied
    # Explicit deletion only after confirmation; version bytes are purged, audit retained.
    with transaction.atomic():
        audit(request.user,'Fshirje dokumenti',doc.pk)
        files = [v.file for v in doc.versions.all()]
        doc.comments.all().delete()
        doc.versions.all().delete()
        Protocol.objects.filter(document=doc).update(document=None)
        doc.delete()
        transaction.on_commit(lambda: [file.delete(save=False) for file in files])
    messages.success(request,'Dokumenti dhe skedarët u fshinë.')
    return redirect('record_list',module='documents')

@login_required
def spreadsheet(request, pk):
    version = get_object_or_404(DocumentVersion,pk=pk)
    if not document_allowed(request.user,version.document,'view'): raise PermissionDenied
    try: data = spreadsheet_rows(version)
    except ValidationError as error:
        messages.error(request,str(error))
        return redirect('document',pk=version.document_id)
    if not data:
        messages.info(request,'Parapamja tabelare është për XLSX dhe CSV.')
        return redirect('document',pk=version.document_id)
    headers, rows = data[0], data[1:]
    q = request.GET.get('q','').lower()
    if q: rows = [row for row in rows if any(q in cell.lower() for cell in row)]
    try:
        column = int(request.GET.get('sort','-1'))
        if 0 <= column < len(headers): rows.sort(key=lambda row: row[column] if column < len(row) else '', reverse=request.GET.get('direction')=='desc')
    except ValueError: pass
    if request.GET.get('format'):
        if not document_allowed(request.user,version.document,'download'): raise PermissionDenied
        return export_table(request,version.document.title,headers,rows)
    return render(request,'school/spreadsheet.html',{'version':version,'headers':headers,'rows':rows[:1000],
        'total':len(rows),'can_download':document_allowed(request.user,version.document,'download'),
        'can_edit':document_allowed(request.user,version.document,'edit') and not version.document.year.archived})

@login_required
def edit_spreadsheet(request, pk):
    version = get_object_or_404(DocumentVersion,pk=pk)
    doc = version.document
    if not document_allowed(request.user,doc,'edit') or doc.year.archived: raise PermissionDenied
    try: data = spreadsheet_rows(version)
    except ValidationError as error:
        messages.error(request,str(error))
        return redirect('document',pk=doc.pk)
    if not data or len(data) > 501 or len(data[0]) > 30:
        messages.info(request,'Redaktimi në shfletues mbështet deri 500 rreshta dhe 30 kolona. Për tabela më të mëdha ngarkoni një version të ri.')
        return redirect('document',pk=doc.pk)
    fields = {f'cell_{r}_{c}':forms.CharField(label=header,required=False,strip=False,max_length=4096,
        initial=data[r+1][c] if c < len(data[r+1]) else '')
        for r in range(len(data)-1) for c,header in enumerate(data[0])}
    EditTable = type('EditTable',(forms.Form,),fields)
    form = EditTable(request.POST or None)
    if request.method == 'POST' and form.is_valid():
        with transaction.atomic():
            Document.objects.select_for_update().get(pk=doc.pk)
            if doc.versions.first().pk != version.pk:
                form.add_error(None,'Ekziston një version më i ri. Hapeni para ndryshimit.')
            else:
                book = Workbook()
                sheet = book.active
                sheet.append([safe_cell(cell) for cell in data[0]])
                for r in range(len(data)-1): sheet.append([safe_cell(form.cleaned_data[f'cell_{r}_{c}']) for c in range(len(data[0]))])
                output = io.BytesIO()
                book.save(output)
                from django.core.files.uploadedfile import SimpleUploadedFile
                add_version(doc,SimpleUploadedFile('tabela-perditesuar.xlsx',output.getvalue()),request.user)
                doc.status = 'draft'
                doc.save(update_fields=['status'])
                messages.success(request,'Tabela u ruajt si version i ri XLSX. Versioni i mëparshëm u ruajt.')
                return redirect('document',pk=doc.pk)
    return render(request,'school/edit_table.html',{'version':version,'form':form,'headers':data[0],
        'rows':[[form[f'cell_{r}_{c}'] for c in range(len(data[0]))] for r in range(len(data)-1)]})

@access('admin','teacher','homeroom')
def import_academic(request, pk, module):
    if module not in {'grades','attendance'}: raise PermissionDenied
    if not allowed(request.user,module,write=True): raise PermissionDenied
    version = get_object_or_404(DocumentVersion,pk=pk)
    doc = version.document
    if not document_allowed(request.user,doc,'download') or doc.year.archived: raise PermissionDenied
    required = ['numri_regjistrit','klasa','paralelja','lenda','periudha','lloji','data','nota'] if module == 'grades' else ['numri_regjistrit','klasa','paralelja','data','statusi','arsyetimi']
    errors, records, seen = [], [], set()
    try: data = spreadsheet_rows(version)
    except ValidationError as error:
        messages.error(request,str(error))
        return redirect('document',pk=doc.pk)
    if ImportBatch.objects.filter(checksum=version.sha256,module=module,year=doc.year).exists(): errors.append('Ky skedar është importuar tashmë në këtë modul dhe vit.')
    if not data or data[0] != required: errors.append('Kolonat e detyrueshme: '+', '.join(required))
    else:
        for index,row in enumerate(data[1:],2):
            try:
                if len(row)!=len(required): raise ValueError('Numër i gabuar kolonash.')
                values = dict(zip(required,row))
                date = forms.DateField().clean(values['data'][:10])
                enrollment = scoped(request.user,'enrollments').get(student__registration_number=values['numri_regjistrit'],
                    classroom__year=doc.year,classroom__level=int(values['klasa']),classroom__parallel=values['paralelja'],starts__lte=date)
                if enrollment.ends and enrollment.ends < date: raise ValueError('Regjistrimi nuk është aktiv në këtë datë.')
                if module == 'grades':
                    assignment = assignments_for(request.user,write=True).get(classroom=enrollment.classroom,subject__name=values['lenda'])
                    period = Period.objects.get(year=doc.year,name=values['periudha'])
                    kind = AssessmentType.objects.get(year=doc.year,name=values['lloji'])
                    record = Grade(enrollment=enrollment,assignment=assignment,period=period,assessment_type=kind,
                        value=int(values['nota']),weight=kind.weight,date=date,author=request.user)
                    key = (str(enrollment.pk),str(assignment.pk),str(kind.pk),str(date),record.value)
                    if Grade.objects.filter(enrollment=enrollment,assignment=assignment,assessment_type=kind,date=date,value=record.value).exists():
                        raise ValueError('Vlerësimi ekziston tashmë.')
                else:
                    record = Attendance(enrollment=enrollment,date=date,status=values['statusi'],reason=values['arsyetimi'],author=request.user)
                    key = (str(enrollment.pk),str(date))
                if key in seen: raise ValueError('Rresht i dyfishuar.')
                seen.add(key)
                record.full_clean()
                records.append(record)
            except (ValueError,ValidationError,ObjectDoesNotExist,MultipleObjectsReturned) as error:
                errors.append(f'Rreshti {index}: të dhëna të pavlefshme ose jashtë autorizimit ({error}).')
    if len(records)>500: errors.append('Kufiri: 500 regjistrime për importim.')
    if request.method=='POST' and request.POST.get('confirm')=='yes' and not errors and records:
        try:
            with transaction.atomic():
                ImportBatch.objects.create(checksum=version.sha256,module=module,year=doc.year,actor=request.user,count=len(records))
                for record in records:
                    if module=='grades': save_grade(record,request.user)
                    else: record.save()
                audit(request.user,'Importim: '+module,doc.pk,{'count':len(records)})
            messages.success(request,f'U importuan {len(records)} regjistrime.')
            return redirect('record_list',module=module)
        except (IntegrityError,ValidationError): errors.append('Importimi u anulua plotësisht për shkak të konfliktit të të dhënave.')
    return render(request,'school/import_academic.html',{'errors':errors,'records':records,'title':MODULES[module][0],'version':version})

@access('admin','secretary')
def import_students(request, pk):
    version = get_object_or_404(DocumentVersion,pk=pk)
    if not document_allowed(request.user,version.document,'download'): raise PermissionDenied
    doc = version.document
    if doc.year.archived: raise PermissionDenied
    try: data = spreadsheet_rows(version)
    except ValidationError as error:
        messages.error(request,str(error))
        return redirect('document',pk=doc.pk)
    required = ['numri_regjistrit','emri','mbiemri','klasa','paralelja','data_regjistrimit']
    errors, preview = [], []
    if not data or data[0] != required: errors.append('Kolonat e detyrueshme, në rend: ' + ', '.join(required))
    else:
        seen = set()
        for index, row in enumerate(data[1:],2):
            try:
                if len(row) != len(required): raise ValueError('Numër i gabuar kolonash.')
                registration, first, last, level, parallel, start = row
                if registration in seen or Student.objects.filter(registration_number=registration).exists(): raise ValueError('Numër regjistri i dyfishuar.')
                seen.add(registration)
                classroom = Classroom.objects.get(year=doc.year,level=int(level),parallel=parallel)
                date = forms.DateField().clean(start[:10])
                student = Student(first_name=first,last_name=last,registration_number=registration)
                student.full_clean()
                preview.append((student,classroom,date))
            except (ValueError, ValidationError, Classroom.DoesNotExist) as error: errors.append(f'Rreshti {index}: {error}')
    if len(preview) > 500: errors.append('Importoni jo më shumë se 500 nxënës për herë.')
    if request.method == 'POST' and not errors and preview and request.POST.get('confirm') == 'yes':
        try:
            with transaction.atomic():
                for student,classroom,date in preview:
                    student.save()
                    enrollment = Enrollment(student=student,classroom=classroom,starts=date)
                    enrollment.full_clean()
                    enrollment.save()
                audit(request.user,'Importim nxënësish',doc.pk,{'count':len(preview),'version':version.number})
            messages.success(request,f'U importuan {len(preview)} nxënës.')
            return redirect('record_list',module='students')
        except (IntegrityError, ValidationError) as error: errors.append('Importimi u anulua plotësisht: konflikt të dhënash.')
    return render(request,'school/import.html',{'version':version,'errors':errors,'preview':preview})

@access('admin','director','pedagogue','teacher','homeroom')
def success_report(request):
    if not allowed(request.user,'grades'): raise PermissionDenied
    periods = Period.objects.select_related('year').order_by('-starts')
    period = get_object_or_404(periods,pk=request.GET['period']) if request.GET.get('period') else periods.first()
    rows, previous = [], None
    report_label, date_range = str(period) if period else '', None
    if period:
        report_type = request.GET.get('report_type','period')
        if report_type not in {'period','month','half','year','range'}: raise PermissionDenied
        if report_type != 'period':
            periods_in_year = list(periods.filter(year=period.year).order_by('starts'))
            first_date, last_date = periods_in_year[0].starts, max(p.ends for p in periods_in_year)
            try:
                if report_type == 'year': start,end = first_date,last_date
                elif report_type == 'month':
                    import calendar
                    start = forms.DateField().clean(request.GET.get('month','')+'-01')
                    end = start.replace(day=calendar.monthrange(start.year,start.month)[1])
                elif report_type == 'half':
                    from datetime import date
                    month_index = first_date.month-1+6
                    midpoint = date(first_date.year+month_index//12,month_index%12+1,1)
                    start,end = (first_date,min(last_date,midpoint-timedelta(days=1))) if request.GET.get('half','1')=='1' else (midpoint,last_date)
                else:
                    start = forms.DateField().clean(request.GET.get('from'))
                    end = forms.DateField().clean(request.GET.get('to'))
                if start>end or end<first_date or start>last_date: raise ValidationError('Intervali nuk i përket vitit shkollor.')
                date_range = max(first_date,start),min(last_date,end)
                report_label = f'{period.year} / {date_range[0]} – {date_range[1]}'
            except ValidationError as error:
                messages.error(request,str(error))
                return render(request,'school/success.html',{'title':'Suksesi dhe statistikat e klasave','periods':periods,
                    'period':period,'rows':[],'classrooms':classes_for(request.user),'subjects':Subject.objects.all()})
        classrooms = classes_for(request.user).filter(year=period.year)
        if request.GET.get('classroom'): classrooms = classrooms.filter(pk=request.GET['classroom'])
        subject = request.GET.get('subject')
        # Teachers may only request their assigned subjects; homeroom teachers can see their class.
        if request.user.role in TEACHERS and subject:
            subject_allowed = assignments_for(request.user).filter(subject_id=subject).exists()
            if not subject_allowed: raise PermissionDenied
        assignment_ids = set(assignments_for(request.user).values_list('pk',flat=True)) if request.user.role in TEACHERS else None
        rows = success_rows(classrooms,period,subject,assignment_ids,date_range)
        previous = periods.filter(year=period.year,ends__lt=period.starts).first() if not date_range else None
        if previous:
            earlier = {r['classroom'].pk:r for r in success_rows(classrooms,previous,subject,assignment_ids)}
            for row in rows:
                old = earlier[row['classroom'].pk]['average']
                row['difference'] = row['average'] - old if row['average'] is not None and old is not None else None
        if request.GET.get('format'):
            headers = ['Klasa','Nxënës','Vlerësime të plota','Të paplota','Mesatarja','Përfshirë','Lëndët','Kategoritë']
            data = [[str(r['classroom']),r['total'],r['complete'],r['incomplete'],r['average'] if r['average'] is not None else '—',r['complete'],
                     '; '.join(f"{s}: {v['average'] if v['average'] is not None else '—'} (n={v['count']})" for s,v in r['subjects'].items()),
                     '; '.join(f'{s}: {v}' for s,v in r['categories'].items())] for r in rows]
            return export_table(request,f'Suksesi dhe statistikat – {report_label}',headers,data)
    return render(request,'school/success.html',{'title':'Suksesi dhe statistikat e klasave','periods':periods,'period':period,
        'rows':rows,'previous':previous,'report_label':report_label,'classrooms':classes_for(request.user),'subjects':Subject.objects.all()})

@login_required
def notifications(request):
    if request.method == 'POST':
        Notification.objects.filter(recipient=request.user,pk=request.POST.get('id')).update(read=True)
        return redirect('notifications')
    return render(request,'school/notifications.html',{'notifications':Notification.objects.filter(recipient=request.user)})

@access('admin','director')
def announce(request):
    class Announcement(forms.Form):
        title = forms.CharField(label='Titulli',max_length=200)
        body = forms.CharField(label='Përmbajtja',widget=forms.Textarea)
        roles = forms.MultipleChoiceField(label='Rolet e marrësve',choices=ROLES,required=False)
        users = forms.ModelMultipleChoiceField(label='Marrës individualë',queryset=User.objects.filter(status='active'),required=False)
        def clean(self):
            values = super().clean()
            if not values.get('roles') and not values.get('users'): raise ValidationError('Zgjidhni marrësit.')
            return values
    form = Announcement(request.POST or None)
    if request.method == 'POST' and form.is_valid():
        recipients = User.objects.filter(Q(role__in=form.cleaned_data['roles']) | Q(pk__in=form.cleaned_data['users'].values_list('pk',flat=True)),status='active').distinct()
        with transaction.atomic():
            Notification.objects.bulk_create([Notification(recipient=u,title=form.cleaned_data['title'],body=form.cleaned_data['body']) for u in recipients])
            audit(request.user,'Njoftim i brendshëm','notifications',{'recipients':recipients.count()})
        messages.success(request,'Njoftimi iu dërgua marrësve të zgjedhur.')
        return redirect('notifications')
    return render(request,'school/form.html',{'form':form,'title':'Njoftim i ri'})

def safe_cell(value):
    if isinstance(value,(int,float,Decimal)): return value
    text = str(value) if value is not None else ''
    return "'" + text if text.startswith(('=','+','-','@','\t','\r')) else text

def export_table(request, title, headers, rows):
    fmt = request.GET.get('format','csv')
    school = SchoolSettings.objects.first()
    name = school.name if school else 'Shkolla fillore Rilindja – Llojan'
    stamp = timezone.localtime().strftime('%d.%m.%Y %H:%M')
    if fmt == 'csv':
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = 'attachment; filename="raporti.csv"'
        response.write('\ufeff')
        writer = csv.writer(response)
        writer.writerow([safe_cell(name),safe_cell(title),stamp,str(request.user)])
        writer.writerow(headers)
        writer.writerows([[safe_cell(c) for c in row] for row in rows])
    elif fmt == 'xlsx':
        book = Workbook()
        sheet = book.active
        sheet.title = 'Raporti'
        sheet.append([safe_cell(name),safe_cell(title),stamp,str(request.user)])
        sheet.append(headers)
        for row in rows: sheet.append([safe_cell(c) for c in row])
        sheet.freeze_panes = 'A3'
        sheet.auto_filter.ref = f'A2:{sheet.cell(2,len(headers)).column_letter}{sheet.max_row}'
        output = io.BytesIO()
        book.save(output)
        response = HttpResponse(output.getvalue(),content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="raporti.xlsx"'
    elif fmt == 'pdf':
        font = Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')
        fontname = 'Helvetica'
        if font.exists():
            if 'DejaVu' not in pdfmetrics.getRegisteredFontNames(): pdfmetrics.registerFont(TTFont('DejaVu',str(font)))
            fontname = 'DejaVu'
        output = io.BytesIO()
        doc = SimpleDocTemplate(output,pagesize=landscape(A4),rightMargin=20,leftMargin=20)
        styles = getSampleStyleSheet()
        for style in styles.byName.values(): style.fontName = fontname
        styles['BodyText'].fontSize = 7
        story = [Paragraph(escape(name),styles['Title']),Paragraph(escape(title),styles['Heading2']),
                 Paragraph(escape(f'{stamp} • {request.user}'),styles['BodyText']),Spacer(1,14)]
        if school and school.logo:
            from reportlab.platypus import Image
            story.insert(0,Image(school.logo.path,width=50,height=50,kind='proportional'))
        cells = [[Paragraph(escape(str(c)),styles['BodyText']) for c in row] for row in [headers] + rows]
        table = Table(cells,repeatRows=1,colWidths=[doc.width/len(headers)]*len(headers))
        table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e8eef6')),
            ('GRID',(0,0),(-1,-1),0.3,colors.HexColor('#ccd6e1')),('VALIGN',(0,0),(-1,-1),'TOP')]))
        story.append(table)
        doc.build(story)
        response = HttpResponse(output.getvalue(),content_type='application/pdf')
        response['Content-Disposition'] = 'attachment; filename="raporti.pdf"'
    else: raise PermissionDenied
    audit(request.user,'Eksport raporti',title,{'format':fmt,'rows':len(rows)})
    response['Cache-Control'] = 'no-store'
    return response
