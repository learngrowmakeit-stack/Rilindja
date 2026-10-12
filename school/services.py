import hashlib
import io
import zipfile
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from django.core.exceptions import ValidationError
from django.db import transaction
from django.db.models import Max
from openpyxl import load_workbook
from .models import *

def audit(user, action, target, details=None):
    Audit.objects.create(actor=user if user and user.is_authenticated else None,
                         action=action, target=str(target), details=details or {})

def grade_snapshot(grade):
    return {'value': grade.value, 'weight': str(grade.weight), 'date': str(grade.date),
            'period': str(grade.period_id), 'type': str(grade.assessment_type_id),
            'enrollment': str(grade.enrollment_id), 'assignment': str(grade.assignment_id), 'revision': grade.revision}

@transaction.atomic
def save_grade(grade, actor, expected_revision=None):
    previous = None
    if Grade.objects.filter(pk=grade.pk).exists():
        old = Grade.objects.select_for_update().get(pk=grade.pk)
        if expected_revision != old.revision: raise ValidationError('Nota është ndryshuar nga dikush tjetër. Rifreskoni faqen.')
        previous = grade_snapshot(old)
        # Keep original weighting; changes to rules only affect future assessments.
        grade.weight = old.weight if old.assessment_type_id == grade.assessment_type_id else grade.assessment_type.weight
        grade.revision = old.revision + 1
    else:
        grade.weight = grade.assessment_type.weight
    if grade.period.year.archived: raise ValidationError('Viti është arkivuar.')
    grade.author = actor
    grade.full_clean()
    grade.save()
    GradeHistory.objects.create(grade=grade, actor=actor, previous=previous, current=grade_snapshot(grade))
    audit(actor, 'Ndryshim note', grade.pk, {'revision': grade.revision})

def rounded(value):
    return value.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP) if value is not None else None

def category(value, year):
    if value is None: return 'Të paplota'
    if value < year.pass_threshold: return 'Pamjaftueshëm'
    if value < Decimal('2.50'): return 'Mjaftueshëm'
    if value < Decimal('3.50'): return 'Mirë'
    if value < Decimal('4.50'): return 'Shumë mirë'
    return 'Shkëlqyeshëm'

def success_rows(classrooms, period, subject_id=None, assignment_ids=None, date_range=None):
    rows = []
    for classroom in classrooms:
        assignments = list(Assignment.objects.filter(classroom=classroom).select_related('subject'))
        if assignment_ids is not None: assignments = [a for a in assignments if a.pk in assignment_ids]
        if subject_id: assignments = [a for a in assignments if str(a.subject_id) == str(subject_id)]
        starts, ends = date_range or (period.starts,period.ends)
        students = Student.objects.filter(Q(enrollment__ends__isnull=True) | Q(enrollment__ends__gte=starts),
            enrollment__classroom=classroom, enrollment__starts__lte=ends).distinct()
        grades = Grade.objects.filter(assignment__in=assignments).select_related('enrollment')
        grades = grades.filter(period__year=period.year,date__gte=starts,date__lte=ends) if date_range else grades.filter(period=period)
        buckets = {}
        for grade in grades:
            key = (grade.enrollment.student_id, grade.assignment_id)
            total, weight = buckets.get(key, (Decimal(0), Decimal(0)))
            buckets[key] = (total + grade.value * grade.weight, weight + grade.weight)
        individuals, subject_means, counts = [], {}, {}
        for assignment in assignments:
            vals = [total / weight for (student, aid), (total, weight) in buckets.items() if aid == assignment.id]
            subject_means[assignment.subject.name] = {'average': rounded(sum(vals) / len(vals)) if vals else None, 'count': len(vals)}
        for student in students:
            values = []
            per_subject = {}
            for assignment in assignments:
                pair = buckets.get((student.id, assignment.id))
                val = pair[0] / pair[1] if pair else None
                per_subject[assignment.subject.name] = rounded(val)
                if val is not None: values.append(val)
            complete = bool(assignments) and len(values) == len(assignments)
            average = sum(values) / len(values) if complete else None
            label = category(average, classroom.year)
            counts[label] = counts.get(label, 0) + 1
            individuals.append({'student': student, 'average': rounded(average), 'raw_average': average,
                                'complete': complete, 'subjects': per_subject, 'category': label})
        included = [i['raw_average'] for i in individuals if i['complete']]
        rows.append({'classroom': classroom, 'total': len(individuals), 'complete': len(included),
                     'incomplete': len(individuals) - len(included), 'subjects': subject_means, 'categories': counts,
                     'average': rounded(sum(included) / len(included)) if included else None, 'individuals': individuals})
    return rows

def validate_upload(upload):
    if upload.size > 20 * 1024 * 1024: raise ValidationError('Madhësia maksimale është 20 MB.')
    ext = Path(upload.name).suffix.lower()
    if ext not in {'.pdf', '.docx', '.xlsx', '.csv'}: raise ValidationError('Lejohen PDF, DOCX, XLSX dhe CSV.')
    data = upload.read()
    upload.seek(0)
    if ext == '.pdf' and not data.startswith(b'%PDF-'): raise ValidationError('PDF i pavlefshëm.')
    if ext in {'.xlsx', '.docx'}:
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                entries = archive.infolist()
                if len(entries) > 3000 or sum(e.file_size for e in entries) > 60 * 1024 * 1024:
                    raise ValidationError('Arkiv shumë i madh pas hapjes.')
                names = archive.namelist()
                required = 'xl/workbook.xml' if ext == '.xlsx' else 'word/document.xml'
                if required not in names or any('vbaProject' in name for name in names):
                    raise ValidationError('Dokument i pavlefshëm ose me makro.')
        except zipfile.BadZipFile: raise ValidationError('Skedar Office i pavlefshëm.')
    if ext == '.csv':
        try: data.decode('utf-8-sig')
        except UnicodeDecodeError: raise ValidationError('CSV duhet të jetë UTF-8.')
        if b'\x00' in data: raise ValidationError('CSV i pavlefshëm.')
    return hashlib.sha256(data).hexdigest()

@transaction.atomic
def add_version(document, upload, actor):
    checksum = validate_upload(upload)
    Document.objects.select_for_update().get(pk=document.pk)
    number = (document.versions.aggregate(n=Max('number'))['n'] or 0) + 1
    version = DocumentVersion.objects.create(document=document, number=number, file=upload,
        original_name=Path(upload.name).name, sha256=checksum, uploaded_by=actor)
    audit(actor, 'Ngarkim dokumenti', document.pk, {'version': number})
    return version

def spreadsheet_rows(version):
    if not version.original_name.lower().endswith(('.xlsx', '.csv')): return []
    with version.file.open('rb') as handle:
        if version.original_name.lower().endswith('.csv'):
            import csv
            data = list(csv.reader(io.StringIO(handle.read().decode('utf-8-sig'))))
        else:
            workbook = load_workbook(handle, read_only=True, data_only=True, keep_links=False)
            sheet = workbook.active
            if sheet.max_row > 10001 or sheet.max_column > 100: raise ValidationError('Kufiri: 10 000 rreshta dhe 100 kolona.')
            data = [list(row) for row in sheet.iter_rows(values_only=True)]
            workbook.close()
    if len(data) > 10001 or any(len(row) > 100 for row in data): raise ValidationError('Tabela tejkalon kufirin.')
    return [[str(cell) if cell is not None else '' for cell in row] for row in data]

@transaction.atomic
def transfer(enrollment, destination, date, actor):
    original = Enrollment.objects.select_for_update().get(pk=enrollment.pk)
    if original.ends or original.classroom.year.archived or destination.year.archived:
        raise ValidationError('Regjistrimi është mbyllur ose viti është arkivuar.')
    if destination.year_id != original.classroom.year_id or destination.pk == original.classroom_id:
        raise ValidationError('Transferimi bëhet në një klasë tjetër të të njëjtit vit.')
    if date <= original.starts: raise ValidationError('Data duhet të jetë pas regjistrimit.')
    if Grade.objects.filter(enrollment=original, date__gte=date).exists() or Attendance.objects.filter(enrollment=original, date__gte=date).exists():
        raise ValidationError('Ekzistojnë nota ose vijueshmëri pas datës së transferimit.')
    from datetime import timedelta
    original.ends = date - timedelta(days=1)
    original.full_clean()
    original.save()
    new = Enrollment(student=original.student, classroom=destination, starts=date)
    new.full_clean()
    new.save()
    audit(actor, 'Transferim nxënësi', original.student_id, {'from': str(original.classroom_id), 'to': str(destination.pk)})
    return new
